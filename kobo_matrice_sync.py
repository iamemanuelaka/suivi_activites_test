#!/usr/bin/env python3
"""
Synchronisation KoboToolbox -> MATRICE DE SUIVI DES ACTIVITES -> envoi par e-mail.

Étapes :
  1. Récupère les soumissions du formulaire via l'API KoboToolbox (v2), photos comprises.
  2. Détermine, pour chaque lot, l'état courant (dernière visite valide).
  3. Met à jour la dernière matrice de MATRIX_DIR (onglet unique « MATRICE DE SUIVI ») et l'enregistre sous
     MATRICE_DE_SUIVI_DES_ACTIVITES_du_JJ-MM-AA.xlsx.
       - jamais touché : colonnes financières (MONTANT, MONTANT DECAISSE, TAUX D'EXECUTION FINANCIERE),
         formules (FIN DES TRAVAUX, moyenne), numéros de lot, régions, entreprises ;
       - début des travaux / délai : renseignés seulement s'ils sont vides dans la matrice.
  4. Tient un journal des visites (historique/visites.csv) : audit, liens photos, position GPS.
  5. Envoie la matrice par e-mail avec un résumé de la semaine.

Usage :
  python kobo_matrice_sync.py                      # cycle complet (API + mail)
  python kobo_matrice_sync.py --no-mail            # sans envoi d'e-mail
  python kobo_matrice_sync.py --skip-if-unchanged  # ne fait rien s'il n'y a aucune nouvelle réponse
  python kobo_matrice_sync.py --from-json data.json --no-mail --no-photos   # test hors ligne
"""
import argparse
import csv
import datetime as dt
import io
import json
import logging
import os
import re
import smtplib
import ssl
import sys
from copy import copy
from email.message import EmailMessage
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment
from openpyxl.utils.units import pixels_to_EMU

# --------------------------------------------------------------------------
# Libellés : valeurs Kobo (noms de choix) -> texte écrit dans la matrice
# --------------------------------------------------------------------------
STATUT = {
    "non_demarre": "Non démarré",
    "demarrage_autorise": "Démarrage autorisé",
    "os_a_reprendre": "Ordre de service à reprendre (erreur sur l'OS)",
    "demarrage_effectif": "Démarrage effectif",
    "travaux_en_cours": "Travaux en cours",
    "travaux_acheves": "Travaux achevés",
    "suspendu": "Travaux suspendus",
    "abandonne": "Abandonné",
}
PGES = {
    "valide": "Validé",
    "transmis_validation": "Transmis en cours de validation",
    "non_transmis": "Non transmis",
    "pas_de_pges": "Pas de PGES",
}
STRUCTURE = {  # liste « rm3kw40 » du formulaire
    "pcr": "PCR",
    "missions_de_contr_le": "Missions de contrôle",
    "consultant": "Consultant",
    "autre": "Autre",
}
STATUTS_A_SURVEILLER = {"suspendu", "abandonne", "os_a_reprendre"}

# Noms des questions du formulaire
F_DEBUT = "Date_de_d_but_des_travaux"
F_DELAI = "delai"
F_RESULTATS = "R_sultats"
F_ONG_OUINON = "Il_y_a_t_il_une_ONG_mobilis_e"
F_STRUCTURE = "Structure"
F_STRUCTURE_AUTRE = "Autre_pr_ciser"


# Intitulés de colonnes (comparaison après normalisation : majuscules, sans accents).
# Les colonnes financières ne figurent volontairement pas ici : le script ne les lit ni ne les écrit.
HEADERS = {
    "region": lambda h: h.startswith("REGION"),
    "departement": lambda h: h.startswith("DEPARTEMENT"),
    "entreprise": lambda h: h.startswith("ENTREPRISE"),
    "debut": lambda h: h.startswith("DEBUT DES TRAVAUX"),
    "delai": lambda h: h.startswith("DELAI"),
    "remise": lambda h: h.startswith("DATE DE REMISE"),
    "resultats": lambda h: h.startswith("RESULTAT"),
    "taux": lambda h: h.startswith("TAUX D'EXECUTION") and "FIN" not in h,
    "obs_tech": lambda h: h.startswith("OBSERVATIONS TECHNIQUES"),
    "pges": lambda h: h.startswith("OBSERVATIONS ENVIRONNEMENTALES"),
    "ong": lambda h: h.startswith("ONG MOBILISEES"),
    "recep_prov": lambda h: h.startswith("DATE DE RECEPTION PROVISOIRE"),
    "recep_def": lambda h: h.startswith("DATE DE RECEPTION DEFINITIVE"),
    "photo": lambda h: h.startswith("PHOTO"),
    "repondant": lambda h: h.startswith("NOM") and "REPONDANT" in h,
}
DATE_FMT = "mm-dd-yy"  # format « date courte » d'Excel (affichage selon la langue du poste), comme le reste de la matrice
SHEET_MATRIX = "MATRICE DE SUIVI"
MATRIX_PATTERN = "MATRICE_DE_SUIVI_DES_ACTIVITES_du_*.xlsx"
PHOTO_COL_WIDTH = 34
THUMB_STORE = (240, 180)   # taille des miniatures conservées dans le dépôt
THUMB_SHOW = (112, 84)     # taille d'affichage dans la cellule (pixels)
ROW_MIN_HEIGHT = 72        # hauteur mini d'une ligne avec photos (points)
TEST_NAMES = {"INIT", "TEST"}  # soumissions techniques (ex. première soumission exigée par Kobo) : ignorées

log = logging.getLogger("kobo_sync")


# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------
def load_env(path: Path) -> None:
    """Charge un fichier .env simple (CLE=valeur) sans écraser l'environnement."""
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def cfg(name: str, default: str = "", required: bool = False) -> str:
    val = (os.environ.get(name) or default).strip()  # une variable vide = valeur par défaut
    if required and not val:
        sys.exit(f"Variable de configuration manquante : {name} (voir .env.example)")
    return val


# --------------------------------------------------------------------------
# 1. Kobo
# --------------------------------------------------------------------------
def kobo_headers() -> dict:
    return {"Authorization": f"Token {cfg('KOBO_TOKEN', required=True)}"}


def fetch_submissions() -> list[dict]:
    base = cfg("KOBO_URL", "https://kf.kobotoolbox.org").rstrip("/")
    uid = cfg("KOBO_ASSET_UID", required=True)
    url = f"{base}/api/v2/assets/{uid}/data/?format=json&limit=1000"
    out = []
    while url:
        r = requests.get(url, headers=kobo_headers(), timeout=60)
        r.raise_for_status()
        payload = r.json()
        out.extend(payload.get("results", []))
        url = payload.get("next")
    log.info("%d soumission(s) récupérée(s) depuis Kobo", len(out))
    return out


def flatten(sub: dict) -> dict:
    """Retire les préfixes de groupes ('avancement/taux_execution' -> 'taux_execution')."""
    flat = {k.split("/")[-1]: v for k, v in sub.items()}
    for k in ("_id", "_submission_time", "_validation_status", "_attachments"):
        flat[k] = sub.get(k)
    return flat


def pdate(value) -> dt.date | None:
    if not value:
        return None
    try:
        return dt.date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def num(value):
    """'35' -> 35 ; '12.5' -> 12.5 ; vide/invalide -> None."""
    if value in (None, ""):
        return None
    try:
        f = float(str(value).replace(",", ".").replace(" ", ""))
    except ValueError:
        return None
    return int(f) if f == int(f) else f


def is_ignored(sub: dict) -> bool:
    """Soumission rejetée dans Kobo (« Not approved ») ou soumission technique de test."""
    vs = sub.get("_validation_status") or {}
    uid = vs.get("uid") if isinstance(vs, dict) else vs
    if uid == "validation_status_not_approved":
        return True
    return str(sub.get("nom_agent") or "").strip().upper() in TEST_NAMES


def lot_number(sub: dict) -> int | None:
    try:
        return int(str(sub.get("id_projet", "")).strip())
    except ValueError:
        return None


def start_tod(s: dict) -> int:
    """Heure d'ouverture du formulaire (métadonnée « start ») en secondes depuis minuit ; sert de départage."""
    m = re.search(r"T(\d{2}):(\d{2}):(\d{2})", str(s.get("start") or ""))
    return int(m[1]) * 3600 + int(m[2]) * 60 + int(m[3]) if m else 0


def sort_key(s: dict) -> tuple:
    """Même règle que le formulaire Kobo (champ ordre_suivi) : date de visite la plus récente ; à date égale,
    formulaire ouvert le plus tard ; en dernier recours, date de soumission."""
    return (pdate(s.get("date_mission")) or dt.date.min, start_tod(s), str(s.get("_submission_time") or ""))


def latest_per_lot(subs: list[dict]) -> dict[int, dict]:
    """État courant de chaque lot = dernière visite valide (règle de sort_key). Une valeur vide dans cette visite
    n'est jamais remplacée par une ancienne valeur ; elle n'écrit simplement rien dans la matrice."""
    out: dict[int, dict] = {}
    for s in sorted((x for x in subs if not is_ignored(x)), key=sort_key):
        n = lot_number(s)
        if n is None:
            log.warning("Soumission %s ignorée : id_projet invalide", s.get("_id"))
            continue
        out[n] = s  # tri croissant : la dernière écrase les précédentes
    return out


# ----- photos ----------------------------------------------------------------
def photo_attachments(sub: dict) -> dict[str, dict]:
    """{'photo_1': {...urls...}, 'photo_2': {...}} d'après _attachments."""
    out = {}
    for att in sub.get("_attachments") or []:
        if att.get("is_deleted"):
            continue
        xpath = str(att.get("question_xpath") or "")
        base = str(att.get("media_file_basename") or str(att.get("filename") or "").split("/")[-1])
        for key in ("photo_1", "photo_2"):
            if xpath.split("/")[-1] == key or (sub.get(key) and base == str(sub.get(key))):
                out.setdefault(key, att)
    return out


def download_image(att: dict) -> bytes | None:
    for field in ("download_medium_url", "download_url"):
        url = att.get(field)
        if not url:
            continue
        try:
            r = requests.get(url, headers=kobo_headers(), timeout=60)
            if r.ok and r.content:
                return r.content
            log.warning("Photo non téléchargée (%s) : HTTP %s", field, r.status_code)
        except requests.RequestException as e:
            log.warning("Photo non téléchargée (%s) : %s", field, e)
    return None


def make_thumb(data: bytes) -> bytes:
    from PIL import Image, ImageOps
    img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
    img.thumbnail(THUMB_STORE)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=72, optimize=True)
    return buf.getvalue()


def photo_path(photo_dir: Path, n: int, idx: int) -> Path:
    return photo_dir / f"lot_{n:02d}_{idx}.jpg"


def refresh_photos(latest: dict[int, dict], new_ids: set[str], photo_dir: Path) -> None:
    """Télécharge (miniatures) les photos des lots dont la dernière visite est nouvelle."""
    photo_dir.mkdir(parents=True, exist_ok=True)
    for n, s in latest.items():
        atts = photo_attachments(s)
        for idx in (1, 2):
            path = photo_path(photo_dir, n, idx)
            att = atts.get(f"photo_{idx}")
            if not att:
                if str(s.get("_id")) in new_ids:
                    path.unlink(missing_ok=True)  # la nouvelle visite n'a pas cette photo
                continue
            if str(s.get("_id")) not in new_ids and path.exists():
                continue
            data = download_image(att)
            if data:
                try:
                    path.write_bytes(make_thumb(data))
                except Exception as e:  # image illisible
                    log.warning("Lot %s photo %s illisible : %s", n, idx, e)


# --------------------------------------------------------------------------
# 2. Matrice
# --------------------------------------------------------------------------
def norm(s) -> str:
    s = re.sub(r"\s+", " ", str(s or "")).strip().upper()
    for a, b in zip("ÉÈÊËÀÂÎÏÔÛÙÇ’‘`", "EEEEAAIIOUUC'''"):
        s = s.replace(a, b)
    return s


def find_columns(ws, row_range=range(1, 11)) -> tuple[int, dict[str, int]]:
    for r in row_range:
        found = {}
        for c in range(1, ws.max_column + 1):
            h = norm(ws.cell(r, c).value)
            if not h:
                continue
            for k, test in HEADERS.items():
                if k not in found and test(h):
                    found[k] = c
        if "taux" in found and "region" in found:
            return r, found
    sys.exit(f"En-têtes introuvables dans l'onglet « {ws.title} » (ligne « TAUX D'EXECUTION »).")


def lot_rows(ws, header_row: int, cols: dict[str, int]) -> dict[int, int]:
    num_col = cols["region"] - 1  # la colonne des numéros est juste avant 'REGIONS'
    rows = {}
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(r, num_col).value
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            rows[int(v)] = r
        elif rows:
            break  # fin du tableau
    return rows


def lot_identity(ws, rows: dict[int, int], cols: dict[str, int]) -> dict[int, dict]:
    ident, region = {}, ""
    for n, r in sorted(rows.items()):
        region = ws.cell(r, cols["region"]).value or region  # cellules fusionnées
        ident[n] = {
            "region": str(region).strip(),
            "departement": str(ws.cell(r, cols["departement"]).value or "").strip(),
            "entreprise": str(ws.cell(r, cols["entreprise"]).value or "").strip(),
        }
    return ident


def set_date(cell, d: dt.date) -> None:
    cell.value = d
    cell.number_format = DATE_FMT


def put(ws, r: int, cols: dict, name: str, value, fmt: str | None = None) -> None:
    c = cols.get(name)
    if c is None or value is None:
        return
    cell = ws.cell(r, c)
    cell.value = value
    if fmt:
        cell.number_format = fmt


def fill_if_empty(ws, r: int, cols: dict, name: str, value, is_date: bool = False) -> None:
    """Renseigne une donnée contractuelle seulement si la case est vide (la matrice reste la référence)."""
    c = cols.get(name)
    if c is None or value is None:
        return
    cell = ws.cell(r, c)
    if cell.value in (None, ""):
        if is_date:
            set_date(cell, value)
        else:
            cell.value = value
        return
    cur = cell.value.date() if isinstance(cell.value, dt.datetime) else cell.value
    if cur != value:
        log.warning("Ligne %d : %s déjà renseigné dans la matrice (%s) ≠ formulaire (%s) – matrice conservée.",
                    r, name, cur, value)


def normalize_taux_scale(ws, rows, col) -> int:
    """Ramène les taux saisis en % (0,05 affiché 5 %) sur l'échelle 0-100 de la matrice."""
    changed = 0
    for r in rows.values():
        c = ws.cell(r, col)
        if "%" in str(c.number_format):
            if isinstance(c.value, (int, float)):
                c.value = round(c.value * 100, 2)
            c.number_format = "General"
            changed += 1
    return changed


# ----- libellés ------------------------------------------------------------
def ong_text(s: dict) -> str | None:
    flag = s.get(F_ONG_OUINON)
    if flag == "non":
        return "Non mobilisée"
    if flag == "oui":
        return (s.get("ong_autre_nom") or "").strip() or "ONG mobilisée"
    return None


def pges_text(s: dict) -> str | None:
    st = s.get("pges_statut")
    if not st:
        return None
    txt = PGES.get(st, st)
    det = re.sub(r"\s+", " ", str(s.get("details_pges") or "")).strip()
    return f"{txt} – {det}" if det and st != "valide" else txt


def structure_text(s: dict) -> str:
    st = s.get(F_STRUCTURE)
    if not st:
        return ""
    if st == "autre":
        return (s.get(F_STRUCTURE_AUTRE) or "").strip() or "Autre"
    return STRUCTURE.get(st, str(st).replace("_", " ").title())


def nom_text(s: dict) -> str:
    nom = re.sub(r"\s+", " ", str(s.get("nom_agent") or "")).strip()
    return nom.title() if nom.islower() or nom.isupper() else nom  # « kone ali » -> « Kone Ali »


def repondant_text(s: dict) -> str:
    """« Nom Prénoms - Structure »"""
    nom, st = nom_text(s), structure_text(s)
    return f"{nom} - {st}" if nom and st else (nom or st)


def obs_tech_text(s: dict) -> str | None:
    statut = s.get("statut")
    if not statut:
        return None
    obs = re.sub(r"\s+", " ", str(s.get("obs_techniques") or "")).strip()
    txt = STATUT.get(statut, statut)
    return f"{txt} – {obs}" if obs else txt


# ----- photos dans la matrice -----------------------------------------------
def clear_photo_images(ws, col0: int, first_row0: int, last_row0: int) -> None:
    def keep(im) -> bool:
        a = getattr(im, "anchor", None)
        m = getattr(a, "_from", None)
        return not (m is not None and m.col == col0 and first_row0 <= m.row <= last_row0)
    ws._images = [im for im in ws._images if keep(im)]


def place_photos(ws, r: int, c: int, n: int, photo_dir: Path) -> bool:
    w, h = THUMB_SHOW
    placed, slot = False, 0
    for idx in (1, 2):
        p = photo_path(photo_dir, n, idx)
        if not p.exists():
            continue
        i, slot = slot, slot + 1
        img = XLImage(io.BytesIO(p.read_bytes()))
        img.width, img.height = w, h
        marker = AnchorMarker(col=c - 1, colOff=pixels_to_EMU(4 + i * (w + 4)), row=r - 1, rowOff=pixels_to_EMU(4))
        img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
        ws.add_image(img)
        placed = True
    if placed:
        cur = ws.row_dimensions[r].height or 15
        ws.row_dimensions[r].height = max(cur, ROW_MIN_HEIGHT)
        col = ws.cell(r, c).column_letter
        if (ws.column_dimensions[col].width or 0) < PHOTO_COL_WIDTH:
            ws.column_dimensions[col].width = PHOTO_COL_WIDTH
    return placed


def update_matrix(wb, latest: dict[int, dict], photo_dir: Path) -> dict:
    ws = wb[SHEET_MATRIX] if SHEET_MATRIX in wb.sheetnames else wb.active
    header_row, cols = find_columns(ws)
    rows = lot_rows(ws, header_row, cols)
    if not rows:
        sys.exit("Aucun lot trouvé dans la matrice (numéros de lot attendus à gauche de « REGIONS »).")
    ident = lot_identity(ws, rows, cols)
    first_row, last_row = min(rows.values()), max(rows.values())

    if n_fixed := normalize_taux_scale(ws, rows, cols["taux"]):
        log.info("Échelle des taux harmonisée (0-100) sur %d cellule(s)", n_fixed)
    missing = [k for k in ("debut", "delai", "remise", "resultats", "photo", "repondant") if k not in cols]
    if missing:
        log.warning("Colonnes absentes de la matrice (ignorées) : %s", ", ".join(missing))
    if "photo" in cols:
        clear_photo_images(ws, cols["photo"] - 1, first_row - 1, last_row - 1)

    for n, s in latest.items():
        r = rows.get(n)
        if r is None:
            log.warning("Lot %s absent de la matrice, soumission ignorée", n)
            continue
        put(ws, r, cols, "obs_tech", obs_tech_text(s))
        put(ws, r, cols, "taux", num(s.get("taux_execution")), "General")
        put(ws, r, cols, "resultats", num(s.get(F_RESULTATS)), "General")
        put(ws, r, cols, "pges", pges_text(s))
        put(ws, r, cols, "ong", ong_text(s))
        if d := pdate(s.get("date_remise_site")):
            if "remise" in cols:
                set_date(ws.cell(r, cols["remise"]), d)
        fill_if_empty(ws, r, cols, "debut", pdate(s.get(F_DEBUT)), is_date=True)
        fill_if_empty(ws, r, cols, "delai", num(s.get(F_DELAI)))
        if s.get("recep_prov") == "oui" and (d := pdate(s.get("date_recep_prov"))) and "recep_prov" in cols:
            set_date(ws.cell(r, cols["recep_prov"]), d)
        if s.get("recep_def") == "oui" and (d := pdate(s.get("date_recep_def"))) and "recep_def" in cols:
            set_date(ws.cell(r, cols["recep_def"]), d)
        if "repondant" in cols:
            cell = ws.cell(r, cols["repondant"])
            cell.value = repondant_text(s)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    if "photo" in cols:  # miniatures issues du cache (dossier photos/), pour tous les lots
        for n, r in rows.items():
            place_photos(ws, r, cols["photo"], n, photo_dir)

    # Ligne de total : on ne touche pas à une formule existante (=AVERAGE(...))
    vals = [ws.cell(r, cols["taux"]).value for r in rows.values()]
    vals = [v for v in vals if isinstance(v, (int, float))]
    mean = round(sum(vals) / len(vals), 2) if vals else 0
    for r in range(last_row + 1, last_row + 6):
        if norm(ws.cell(r, cols["region"]).value).startswith("TAUX D"):
            cell = ws.cell(r, cols["taux"])
            if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.value = mean
            break
    return {"rows": rows, "ident": ident, "mean": mean}


# ----- journal des visites (hors matrice) ------------------------------------
JOURNAL_FIELDS = ["id_kobo", "date_visite", "lot", "region", "departement", "entreprise", "statut",
                  "taux_execution", "resultats", "pges", "ong", "nom", "structure", "gps_maps",
                  "photo_1", "photo_2", "commentaires", "soumis_le"]


def journal_ids(path: Path) -> set[str]:
    if not path.exists():
        return set()
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return {row["id_kobo"] for row in csv.DictReader(fh) if row.get("id_kobo")}


def gps_link(value) -> str:
    parts = str(value or "").split()
    return f"https://www.google.com/maps?q={parts[0]},{parts[1]}" if len(parts) >= 2 else ""


def append_journal(path: Path, subs: list[dict], ident: dict[int, dict], known: set[str]) -> int:
    """Une ligne par visite (audit : liens photos Kobo, position GPS). Sans doublon, hors soumissions ignorées."""
    path.parent.mkdir(parents=True, exist_ok=True)
    new = [s for s in sorted(subs, key=sort_key) if not is_ignored(s) and str(s.get("_id")) not in known]
    if not new:
        return 0
    exists = path.exists()
    with path.open("a", encoding="utf-8-sig" if not exists else "utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=JOURNAL_FIELDS)
        if not exists:
            w.writeheader()
        for s in new:
            i = ident.get(lot_number(s), {})
            atts = photo_attachments(s)
            w.writerow({
                "id_kobo": s.get("_id"), "date_visite": pdate(s.get("date_mission")) or "", "lot": lot_number(s),
                "region": i.get("region", ""), "departement": i.get("departement", ""),
                "entreprise": i.get("entreprise", ""), "statut": STATUT.get(s.get("statut"), s.get("statut") or ""),
                "taux_execution": num(s.get("taux_execution")) if num(s.get("taux_execution")) is not None else "",
                "resultats": num(s.get(F_RESULTATS)) if num(s.get(F_RESULTATS)) is not None else "",
                "pges": pges_text(s) or "", "ong": ong_text(s) or "", "nom": nom_text(s),
                "structure": structure_text(s), "gps_maps": gps_link(s.get("gps_chantier")),
                "photo_1": (atts.get("photo_1") or {}).get("download_url", ""),
                "photo_2": (atts.get("photo_2") or {}).get("download_url", ""),
                "commentaires": re.sub(r"\s+", " ", str(s.get("commentaires_libres") or "")).strip(),
                "soumis_le": str(s.get("_submission_time") or "")[:16].replace("T", " "),
            })
    return len(new)


# ----- fichiers --------------------------------------------------------------
def find_base_matrix(matrix_dir: Path, explicit: str | None) -> Path:
    if explicit:
        return Path(explicit)
    files = sorted(matrix_dir.glob(MATRIX_PATTERN), key=matrix_date)
    if not files:
        sys.exit(f"Aucune matrice '{MATRIX_PATTERN}' trouvée dans {matrix_dir}")
    return files[-1]


def matrix_date(p: Path) -> tuple:
    m = re.search(r"du_(\d{2})-(\d{2})-(\d{2})", p.name)
    d = dt.date(2000 + int(m[3]), int(m[2]), int(m[1])) if m else dt.date.min
    return (d, p.stat().st_mtime)


def prune_old(matrix_dir: Path, keep: int, current: Path) -> None:
    """Ne conserve que les `keep` matrices les plus récentes (le dépôt Git garde l'historique)."""
    files = sorted(matrix_dir.glob(MATRIX_PATTERN), key=matrix_date, reverse=True)
    for f in files[keep:]:
        if f != current:
            f.unlink()
            log.info("Ancienne matrice supprimée : %s", f.name)


# --------------------------------------------------------------------------
# 3. Résumé et e-mail
# --------------------------------------------------------------------------
def build_summary(info: dict, latest: dict[int, dict], today: dt.date) -> dict:
    since = today - dt.timedelta(days=7)
    reported, missing, watch = [], [], []
    for n in sorted(info["rows"]):
        s = latest.get(n)
        label = f"Lot {n:02d} – {info['ident'][n]['departement']} ({info['ident'][n]['entreprise']})"
        d = pdate(s.get("date_mission")) if s else None
        if s and d and d >= since:
            reported.append(label)
        else:
            missing.append(label)
        if s and s.get("statut") in STATUTS_A_SURVEILLER:
            watch.append(f"{label} : {STATUT[s['statut']]}")
    return {"reported": reported, "missing": missing, "watch": watch, "mean": info["mean"], "total": len(info["rows"])}


def mail_body(summ: dict, today: dt.date) -> str:
    lines = [
        "Bonjour,",
        "",
        f"Veuillez trouver ci-joint la matrice de suivi des activités mise à jour au {today:%d/%m/%Y}.",
        "",
        f"Taux d'exécution moyen : {summ['mean']} %",
        f"Lots visités ces 7 derniers jours : {len(summ['reported'])} sur {summ['total']}",
    ]
    if summ["missing"]:
        lines += ["", "Lots sans visite cette semaine :"] + [f"  - {x}" for x in summ["missing"]]
    if summ["watch"]:
        lines += ["", "Points d'attention :"] + [f"  - {x}" for x in summ["watch"]]
    lines += ["", "Message généré automatiquement à partir des données KoboToolbox.", "", "Cordialement"]
    return "\n".join(lines)


def send_mail(attachment: Path, summ: dict, today: dt.date) -> None:
    host = cfg("SMTP_HOST", required=True)
    port = int(cfg("SMTP_PORT", "587"))
    user = cfg("SMTP_USER", required=True)
    password = cfg("SMTP_PASSWORD", required=True)
    sender = cfg("MAIL_FROM", user)
    to = [a.strip() for a in cfg("MAIL_TO", required=True).split(",") if a.strip()]

    msg = EmailMessage()
    msg["Subject"] = f"Matrice de suivi des activités – {today:%d/%m/%Y}"
    msg["From"], msg["To"] = sender, ", ".join(to)
    msg.set_content(mail_body(summ, today))
    msg.add_attachment(
        attachment.read_bytes(), maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=attachment.name,
    )
    ctx = ssl.create_default_context()
    if port == 465:
        with smtplib.SMTP_SSL(host, port, context=ctx, timeout=60) as smtp:
            smtp.login(user, password)
            smtp.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=60) as smtp:
            smtp.starttls(context=ctx)
            smtp.login(user, password)
            smtp.send_message(msg)
    log.info("E-mail envoyé à %s", ", ".join(to))


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--from-json", help="Fichier JSON de soumissions (test hors ligne)")
    ap.add_argument("--base", help="Matrice de départ (par défaut : la plus récente de MATRIX_DIR)")
    ap.add_argument("--no-mail", action="store_true", help="Ne pas envoyer d'e-mail")
    ap.add_argument("--skip-if-unchanged", action="store_true",
                    help="Ne rien enregistrer s'il n'y a aucune nouvelle soumission")
    ap.add_argument("--no-photos", action="store_true", help="Ne pas télécharger de nouvelles photos")
    ap.add_argument("--env", default=".env", help="Fichier de configuration (défaut : .env)")
    args = ap.parse_args()

    load_env(Path(args.env))
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                        handlers=[logging.StreamHandler(), logging.FileHandler("sync.log", encoding="utf-8")])

    matrix_dir = Path(cfg("MATRIX_DIR", "./matrices"))
    photo_dir = Path(cfg("PHOTO_DIR", "./photos"))
    journal = Path(cfg("HISTORY_DIR", "./historique")) / "visites.csv"
    matrix_dir.mkdir(parents=True, exist_ok=True)
    today = dt.date.today()

    if args.from_json:
        raw = json.loads(Path(args.from_json).read_text(encoding="utf-8"))
        raw = raw.get("results", raw) if isinstance(raw, dict) else raw
    else:
        raw = fetch_submissions()
    subs = [flatten(s) for s in raw]

    known = journal_ids(journal)
    new_ids = {str(s.get("_id")) for s in subs if not is_ignored(s) and str(s.get("_id")) not in known}
    if args.skip_if_unchanged and not new_ids:
        log.info("Aucune nouvelle soumission : matrice inchangée, rien n'est enregistré.")
        return
    log.info("%d nouvelle(s) soumission(s)", len(new_ids))

    base = find_base_matrix(matrix_dir, args.base)
    log.info("Matrice de départ : %s", base.name)
    wb = load_workbook(base)

    latest = latest_per_lot(subs)
    log.info("%d lot(s) avec au moins une visite valide", len(latest))
    if not args.no_photos:
        refresh_photos(latest, new_ids, photo_dir)

    info = update_matrix(wb, latest, photo_dir)
    added = append_journal(journal, subs, info["ident"], known)
    log.info("%d visite(s) ajoutée(s) au journal %s", added, journal)
    wb.calculation.fullCalcOnLoad = True  # Excel recalcule les formules (FIN DES TRAVAUX, moyenne) à l'ouverture

    out = matrix_dir / f"MATRICE_DE_SUIVI_DES_ACTIVITES_du_{today:%d-%m-%y}.xlsx"
    wb.save(out)
    log.info("Matrice enregistrée : %s (taux moyen %s %%)", out, info["mean"])
    prune_old(matrix_dir, keep=int(cfg("MATRIX_KEEP", "8")), current=out)
    if os.environ.get("GITHUB_OUTPUT"):
        with open(os.environ["GITHUB_OUTPUT"], "a", encoding="utf-8") as fh:
            fh.write(f"matrix_path={out}\n")

    if args.no_mail:
        log.info("Envoi d'e-mail désactivé (--no-mail)")
        return
    send_mail(out, build_summary(info, latest, today), today)


if __name__ == "__main__":
    main()
